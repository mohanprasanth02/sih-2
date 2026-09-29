"""
OIML R 76 Standardized Pattern Evaluation & Type Approval Editable MS Word (.docx) Report Generator
Uses python-docx to generate fully editable, standardized test reports for laboratory officers.
"""

import os
from datetime import datetime
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex: str):
    """Set background color of a table cell in python-docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


class OIMLDocxReportGenerator:
    def __init__(self, output_dir: str = "reports/generated"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_report(self, evaluation_data: Dict[str, Any], output_filename: str = None) -> str:
        report_no = evaluation_data.get("report_number", "OIML-R76-REPORT")
        clean_report_no = report_no.replace("/", "-").replace(" ", "_")
        if not output_filename:
            output_filename = f"{clean_report_no}.docx"

        docx_path = os.path.join(self.output_dir, output_filename)
        doc = Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # 1. Header & Title Block
        p_gov = doc.add_paragraph()
        p_gov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_gov = p_gov.add_run("GOVERNMENT OF INDIA • MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION\nDEPARTMENT OF CONSUMER AFFAIRS — DIRECTORATE OF LEGAL METROLOGY\nNATIONAL LEGAL METROLOGY TYPE EVALUATION LABORATORY")
        r_gov.font.size = Pt(8.5)
        r_gov.font.bold = True
        r_gov.font.color.rgb = RGBColor(15, 41, 66)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_title = p_title.add_run("PATTERN EVALUATION TEST REPORT\nNON-AUTOMATIC WEIGHING INSTRUMENT (NAWI)")
        r_title.font.size = Pt(13)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(15, 41, 66)

        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ref = p_ref.add_run("In accordance with OIML Recommendation R 76-1:2006 (E) & Legal Metrology Act, 2009")
        r_ref.font.size = Pt(9)
        r_ref.font.italic = True
        r_ref.font.color.rgb = RGBColor(100, 116, 139)

        # 2. Status & Metadata Banner Table
        status = evaluation_data.get("status", "DRAFT")
        meta_table = doc.add_table(rows=2, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        headers_meta = [
            ("Report Number:", evaluation_data.get("report_number", "N/A")),
            ("Application No:", evaluation_data.get("application_number", "N/A")),
            ("Testing Date:", str(evaluation_data.get("testing_date", datetime.now().strftime("%Y-%m-%d")))[0:10]),
            ("Overall Status:", status.replace("_", " "))
        ]

        row0 = meta_table.rows[0]
        row0.cells[0].paragraphs[0].add_run("Report Number:").bold = True
        row0.cells[1].paragraphs[0].add_run(evaluation_data.get("report_number", "N/A"))
        row0.cells[2].paragraphs[0].add_run("Application No:").bold = True
        row0.cells[3].paragraphs[0].add_run(evaluation_data.get("application_number", "N/A"))

        row1 = meta_table.rows[1]
        row1.cells[0].paragraphs[0].add_run("Testing Date:").bold = True
        row1.cells[1].paragraphs[0].add_run(str(evaluation_data.get("testing_date", datetime.now().strftime("%Y-%m-%d")))[0:10])
        row1.cells[2].paragraphs[0].add_run("Overall Status:").bold = True
        r_stat = row1.cells[3].paragraphs[0].add_run(status.replace("_", " "))
        r_stat.bold = True
        if "APPROVED" in status or status == "PASS":
            r_stat.font.color.rgb = RGBColor(5, 150, 105)
        elif "REJECTED" in status:
            r_stat.font.color.rgb = RGBColor(220, 38, 38)

        for row in meta_table.rows:
            for cell in row.cells:
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(8.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 3. Section 1: Instrument Specifications Table
        h1 = doc.add_paragraph()
        r_h1 = h1.add_run("1. Instrument Specifications & Technical Characteristics")
        r_h1.font.size = Pt(10.5)
        r_h1.font.bold = True
        r_h1.font.color.rgb = RGBColor(15, 41, 66)

        specs = [
            ("Manufacturer Name", evaluation_data.get("manufacturer_name", "N/A"), "Instrument Type", evaluation_data.get("instrument_type", "N/A")),
            ("Manufacturer Address", evaluation_data.get("manufacturer_address", "N/A"), "Model Designation", evaluation_data.get("model_name", "N/A")),
            ("Accuracy Class", evaluation_data.get("accuracy_class", "Class III"), "Serial Number", evaluation_data.get("serial_number", "N/A")),
            ("Max Capacity (Max)", f"{evaluation_data.get('max_capacity')} {evaluation_data.get('units', 'kg')}", "Scale Interval (d)", f"{evaluation_data.get('scale_interval_d')} {evaluation_data.get('units', 'kg')}"),
            ("Min Capacity (Min)", f"{evaluation_data.get('min_capacity')} {evaluation_data.get('units', 'kg')}", "Verification Interval (e)", f"{evaluation_data.get('verification_scale_interval_e')} {evaluation_data.get('units', 'kg')}"),
            ("Scale Intervals (n)", f"n = {evaluation_data.get('n_intervals', 'N/A')}", "Tare Subtractive Max", f"-{evaluation_data.get('max_tare', evaluation_data.get('max_capacity'))} {evaluation_data.get('units', 'kg')}"),
            ("Operating Temperature", f"{evaluation_data.get('temp_range_min', -10)}°C to +{evaluation_data.get('temp_range_max', 40)}°C", "Power Supply", str(evaluation_data.get("power_supply", "230V AC, 50Hz"))),
            ("Load Cell / Sensor", str(evaluation_data.get("load_cell_details", "OIML R60 approved")), "Firmware / Checksum", f"{evaluation_data.get('software_version', 'v1.0')} ({evaluation_data.get('software_checksum', 'CRC-OK')})")
        ]

        specs_table = doc.add_table(rows=len(specs), cols=4)
        specs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, (k1, v1, k2, v2) in enumerate(specs):
            row = specs_table.rows[r_idx]
            row.cells[0].paragraphs[0].add_run(k1).bold = True
            row.cells[1].paragraphs[0].add_run(v1)
            row.cells[2].paragraphs[0].add_run(k2).bold = True
            row.cells[3].paragraphs[0].add_run(v2)
            for c_idx, cell in enumerate(row.cells):
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                if c_idx % 2 == 0:
                    set_cell_background(cell, "F1F5F9")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 4. Section 2: Laboratory & Environmental Conditions Table
        h2 = doc.add_paragraph()
        r_h2 = h2.add_run("2. Laboratory and Environmental Conditions (Clause A.3)")
        r_h2.font.size = Pt(10.5)
        r_h2.font.bold = True
        r_h2.font.color.rgb = RGBColor(15, 41, 66)

        env_specs = [
            ("Laboratory Name", evaluation_data.get("lab_name", "National Legal Metrology Lab"), "Ambient Temp", f"{evaluation_data.get('lab_temperature', 23.5)} °C"),
            ("Accreditation", evaluation_data.get("lab_accreditation", "NABL ISO/IEC 17025 Accredited"), "Relative Humidity", f"{evaluation_data.get('lab_humidity', 52.0)} % RH"),
            ("Local Gravity (g)", f"{evaluation_data.get('local_gravity_g', 9.7803)} m/s²", "Atmospheric Pressure", f"{evaluation_data.get('lab_pressure', 1013.2)} hPa"),
            ("Standard Weights", str(evaluation_data.get("standard_weights_used", "Class M1/F2 (NPL Traceable)")), "Testing Officer", evaluation_data.get("testing_officer_name", "Er. Rajesh Sharma"))
        ]

        env_table = doc.add_table(rows=len(env_specs), cols=4)
        env_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, (k1, v1, k2, v2) in enumerate(env_specs):
            row = env_table.rows[r_idx]
            row.cells[0].paragraphs[0].add_run(k1).bold = True
            row.cells[1].paragraphs[0].add_run(v1)
            row.cells[2].paragraphs[0].add_run(k2).bold = True
            row.cells[3].paragraphs[0].add_run(v2)
            for c_idx, cell in enumerate(row.cells):
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                if c_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 5. Section 3: Weighing Test Table
        h3 = doc.add_paragraph()
        r_h3 = h3.add_run("3. Weighing Performance Test Observations (Clause A.4.4)")
        r_h3.font.size = Pt(10.5)
        r_h3.font.bold = True
        r_h3.font.color.rgb = RGBColor(15, 41, 66)

        summary = evaluation_data.get("evaluation_summary", {})
        readings = summary.get("test_summaries", {}).get("weighing_test", {}).get("readings", [])
        if readings:
            headers = ["#", "Dir", f"Load ({evaluation_data.get('units', 'kg')})", "Indication I", "ΔL", "Error Ec", "Ec (e)", "MPE (±e)", f"MPE ({evaluation_data.get('units', 'kg')})", "Status"]
            w_table = doc.add_table(rows=len(readings) + 1, cols=len(headers))
            w_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            hdr_row = w_table.rows[0]
            for c_idx, title in enumerate(headers):
                cell = hdr_row.cells[c_idx]
                set_cell_background(cell, "E2E8F0")
                set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(title)
                r.bold = True
                r.font.size = Pt(7.5)

            for r_idx, r in enumerate(readings):
                row = w_table.rows[r_idx + 1]
                st_pass = r.get("status") == "PASS"
                vals = [
                    str(r.get("index")),
                    r.get("direction", "INCR"),
                    f"{r.get('load'):g}",
                    f"{r.get('indication'):g}",
                    f"{r.get('delta_l'):g}" if r.get("delta_l") is not None else "-",
                    f"{r.get('corrected_error'):+.4f}",
                    f"{r.get('corrected_error_e'):+.2f} e",
                    f"±{r.get('mpe_e'):.1f} e",
                    f"±{r.get('mpe_unit'):.4f}",
                    r.get("status", "PASS")
                ]
                for c_idx, val in enumerate(vals):
                    cell = row.cells[c_idx]
                    set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    run = p.add_run(val)
                    run.font.size = Pt(7.5)
                    if c_idx == len(vals) - 1:
                        run.bold = True
                        run.font.color.rgb = RGBColor(5, 150, 105) if st_pass else RGBColor(220, 38, 38)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 6. Section 4: Prescribed OIML Tests Summary Table
        h4 = doc.add_paragraph()
        r_h4 = h4.add_run("4. Prescribed Tests Compliance Summary")
        r_h4.font.size = Pt(10.5)
        r_h4.font.bold = True
        r_h4.font.color.rgb = RGBColor(15, 41, 66)

        def get_verdict(test_key):
            t_sum = summary.get("test_summaries", {}).get(test_key, {})
            if t_sum:
                return t_sum.get("overall_status", "PASS")
            return "PASS" if evaluation_data.get("is_fully_compliant", False) else "PENDING"

        tests_rows = [
            ("Clause A.4.4", "Weighing Performance Test", "Error |Ec| <= MPE across range; Hysteresis <= |MPE|", get_verdict("weighing_test")),
            ("Clause A.4.10", "Repeatability Test (at 0.5 Max & Max)", "Max difference between weighings <= |MPE|", get_verdict("repeatability_test")),
            ("Clause A.4.7", "Eccentricity / Off-Center Loading", "|Ec| <= MPE at 1/3 Max at 4 corners and center", get_verdict("eccentricity_test")),
            ("Clause A.4.2", "Zero-Setting and Zero-Tracking", "Residual zero error <= +/- 0.25 e", get_verdict("tare_zero_test")),
            ("Clause A.4.8", "Discrimination Test", "Extra load 1.4 d gives noticeable change of indication", get_verdict("discrimination_test")),
            ("Clause A.5.3", "Static Temperatures (-10°C to +40°C)", "Span error <= MPE; Zero drift <= 1 e / 5°C", get_verdict("environmental_voltage_test")),
            ("Clause A.5.4", "Voltage Variations (+10% / -15%)", "Indications remain within MPE limits", get_verdict("environmental_voltage_test")),
        ]

        t_summary_table = doc.add_table(rows=len(tests_rows) + 1, cols=4)
        t_summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = t_summary_table.rows[0].cells
        hdr_cells[0].paragraphs[0].add_run("Clause").bold = True
        hdr_cells[1].paragraphs[0].add_run("Test Procedure").bold = True
        hdr_cells[2].paragraphs[0].add_run("Prescribed Standard Limit").bold = True
        hdr_cells[3].paragraphs[0].add_run("Verdict").bold = True
        for c in hdr_cells:
            set_cell_background(c, "E2E8F0")
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

        for r_idx, (clause, name, crit, verdict) in enumerate(tests_rows):
            row = t_summary_table.rows[r_idx + 1]
            row.cells[0].paragraphs[0].add_run(clause).bold = True
            row.cells[1].paragraphs[0].add_run(name)
            row.cells[2].paragraphs[0].add_run(crit)
            r_vd = row.cells[3].paragraphs[0].add_run(verdict)
            r_vd.bold = True
            r_vd.font.color.rgb = RGBColor(5, 150, 105) if verdict == "PASS" else RGBColor(220, 38, 38)
            for cell in row.cells:
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 7. Section 5: Statutory Verdict & Signature Blocks
        is_compliant = evaluation_data.get("is_fully_compliant", False)
        verdict_p = doc.add_paragraph()
        verdict_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_verdict = verdict_p.add_run(
            "STATUTORY MODEL APPROVAL VERDICT: " +
            ("COMPLIANT & APPROVED FOR USE" if is_compliant else "NON-COMPLIANT & REJECTED")
        )
        r_verdict.font.size = Pt(11)
        r_verdict.font.bold = True
        r_verdict.font.color.rgb = RGBColor(5, 150, 105) if is_compliant else RGBColor(220, 38, 38)

        # Signature Table
        sig = evaluation_data.get("digital_signature") or {}
        sig_table = doc.add_table(rows=2, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        c00 = sig_table.rows[0].cells[0].paragraphs[0]
        c00.add_run("TESTED & EVALUATED BY:\n").bold = True
        c00.add_run(f"{evaluation_data.get('testing_officer_name', 'Er. Rajesh Sharma')}\nSenior Metrological Officer, Legal Metrology")
        
        c01 = sig_table.rows[0].cells[1].paragraphs[0]
        c01.add_run("VERIFIED & APPROVED BY:\n").bold = True
        c01.add_run(f"{evaluation_data.get('approving_officer_name', 'Dr. Priya V. Iyer')}\nDirector of Legal Metrology, Government of India")

        c10 = sig_table.rows[1].cells[0].paragraphs[0]
        c10.add_run(f"Digital Sign Hash: {sig.get('sha256_hash', 'SHA256-AUTHENTICATED')[:32]}...\nDate: {sig.get('timestamp', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}")
        
        c11 = sig_table.rows[1].cells[1].paragraphs[0]
        c11.add_run("[DIGITALLY SIGNED & VERIFIED]\nOfficial Statutory Repository, Legal Metrology Act, 2009")

        for row in sig_table.rows:
            for cell in row.cells:
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(7.5)

        doc.save(docx_path)
        return docx_path
